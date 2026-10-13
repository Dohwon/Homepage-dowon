#!/usr/bin/env python3
"""Read all staged documents and Office files for literal section/sheet structure.

This deterministic pass supplements, but does not count as, Codex semantic
extraction. Only titles, headings, sheet names and field names enter the graph.
"""
from __future__ import annotations

import hashlib
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from graphify.detect import DOC_EXTENSIONS, OFFICE_EXTENSIONS, docx_to_markdown

ROOT = Path(__file__).resolve().parent.parent
FULL = ROOT / ".local-graphify" / "full"
STAGE = FULL / "source"
OUTPUT = FULL / "codex-batches" / "0008.json"
STATUS = FULL / "document-index-status.json"
COMMON = {"readme", "overview", "summary", "introduction", "project", "references", "contents", "목차", "개요", "요약", "참고", "프로젝트", "데이터", "결과", "설명", "기능", "구현", "테스트"}
MARKDOWN_HEADING = re.compile(r"^\s{0,3}#{1,4}\s+(.+?)\s*#*\s*$")
HTML_HEADING = re.compile(r"<h[1-4][^>]*>(.*?)</h[1-4]>", re.I | re.S)


def identity(*parts: str) -> str:
    return "doc_" + hashlib.sha256("\0".join(parts).encode()).hexdigest()[:20]


def project_of(relative: str) -> str:
    parts = relative.split("/")
    return "/".join(parts[:2]) if parts[0] == "finish" and len(parts) > 1 else parts[0]


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip(" #*|\t\r\n")[:120]


def headings(text: str, suffix: str) -> list[tuple[str, str]]:
    result = []
    for number, line in enumerate(text.splitlines(), 1):
        match = MARKDOWN_HEADING.match(line)
        if match:
            title = clean(match.group(1))
            if title:
                result.append((title, f"line {number}"))
        elif suffix == ".html":
            for match in HTML_HEADING.finditer(line):
                title = clean(re.sub(r"<[^>]+>", "", match.group(1)))
                if title:
                    result.append((title, f"line {number}"))
        if len(result) >= 12:
            break
    if not result:
        for number, line in enumerate(text.splitlines(), 1):
            title = clean(line)
            if 8 <= len(title) <= 100 and not title.startswith(("{", "[", "import ", "const ", "//", "<!--")):
                result.append((title, f"line {number}; candidate title"))
                break
    return list(dict(((title.casefold(), location), (title, location)) for title, location in result).values())


def main() -> None:
    from openpyxl import load_workbook

    manifest = json.loads((FULL / "scan-manifest.json").read_text(encoding="utf8"))
    nodes, edges = [], []
    by_title = defaultdict(list)
    counts = Counter()
    failures = {}
    sheet_rows_reported = 0
    large_workbooks_metadata_only = 0
    for relative in manifest["staged_paths"]:
        file = STAGE / relative
        suffix = file.suffix.lower()
        if suffix not in DOC_EXTENSIONS | OFFICE_EXTENSIONS:
            continue
        file_id = identity(relative, "file")
        file_nodes = []
        file_edges = []
        try:
            if suffix == ".xlsx":
                if file.stat().st_size > 500_000:
                    with zipfile.ZipFile(file) as archive:
                        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
                    sheets = [item.attrib["name"] for item in workbook_xml.iter() if item.tag.endswith("}sheet")]
                    large_workbooks_metadata_only += 1
                    workbook = None
                else:
                    workbook = load_workbook(file, read_only=True, data_only=True)
                    sheets = [sheet.title for sheet in workbook]
                for sheet_name in sheets:
                    sheet_id = identity(relative, "sheet", sheet_name)
                    file_nodes.append({"id": sheet_id, "label": sheet_name, "file_type": "concept",
                                       "source_file": relative, "source_location": f"sheet {sheet_name}"})
                    file_edges.append({"source": file_id, "target": sheet_id, "relation": "has_sheet",
                                       "confidence": "EXTRACTED", "confidence_score": 1.0,
                                       "source_file": relative, "source_location": f"sheet {sheet_name}"})
                    if workbook is None:
                        continue
                    sheet = workbook[sheet_name]
                    first_row = None
                    sheet_rows_reported += sheet.max_row or 0
                    for row in sheet.iter_rows(min_row=1, max_row=10, max_col=30, values_only=True):
                        if first_row is None and any(value is not None for value in row):
                            first_row = row
                    for column, value in enumerate((first_row or ())[:30], 1):
                        field_name = clean(value) if value is not None else ""
                        if not field_name or len(field_name) < 2:
                            continue
                        field_id = identity(relative, "field", sheet_name, str(column))
                        file_nodes.append({"id": field_id, "label": field_name, "file_type": "concept",
                                           "source_file": relative, "source_location": f"sheet {sheet_name}, column {column}"})
                        file_edges.append({"source": sheet_id, "target": field_id, "relation": "first_row_field",
                                           "confidence": "EXTRACTED", "confidence_score": 1.0,
                                           "source_file": relative, "source_location": f"sheet {sheet_name}, column {column}"})
                if workbook is not None:
                    workbook.close()
            else:
                text = docx_to_markdown(file) if suffix == ".docx" else file.read_text(encoding="utf8", errors="replace")
                for title, location in headings(text, suffix):
                    heading_id = identity(relative, "heading", title, location)
                    file_nodes.append({"id": heading_id, "label": title, "file_type": "concept",
                                       "source_file": relative, "source_location": location})
                    file_edges.append({"source": file_id, "target": heading_id, "relation": "has_section",
                                       "confidence": "EXTRACTED", "confidence_score": 1.0,
                                       "source_file": relative, "source_location": location})
                    key = title.casefold()
                    if len(key) >= 8 and key not in COMMON:
                        by_title[key].append((relative, heading_id))
        except Exception as error:
            failures[relative] = type(error).__name__
            continue
        counts["office" if suffix in OFFICE_EXTENSIONS else "document"] += 1
        nodes.append({"id": file_id, "label": file.name, "file_type": "document",
                      "source_file": relative, "source_location": "document structure only"})
        nodes.extend(file_nodes)
        edges.extend(file_edges)
        if sum(counts.values()) % 250 == 0:
            print(f"Indexed {sum(counts.values())} documents/Office files", flush=True)
    cross_edges = 0
    for matches in by_title.values():
        if len(matches) < 2 or len(matches) > 10:
            continue
        for index, (relative, node_id) in enumerate(matches):
            other = next(((source, other_id) for source, other_id in matches[index + 1:]
                          if project_of(source) != project_of(relative)), None)
            if other is None:
                continue
            edges.append({"source": node_id, "target": other[1], "relation": "same_section_title",
                          "confidence": "EXTRACTED", "confidence_score": 1.0,
                          "source_file": relative, "source_location": "identical section title"})
            cross_edges += 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({"nodes": nodes, "edges": edges, "hyperedges": []}, ensure_ascii=False), encoding="utf8")
    STATUS.write_text(json.dumps({"files_indexed": dict(counts), "failed_files": failures,
                                  "xlsx_rows_reported_by_metadata": sheet_rows_reported,
                                  "large_workbooks_metadata_only": large_workbooks_metadata_only,
                                  "nodes": len(nodes), "edges": len(edges),
                                  "cross_project_exact_titles": cross_edges, "semantic_reviewed": False},
                                 ensure_ascii=False, indent=2) + "\n", encoding="utf8")
    print(json.dumps({"indexed": dict(counts), "failed": len(failures),
                      "large_workbooks_metadata_only": large_workbooks_metadata_only,
                      "xlsx_rows_reported_by_metadata": sheet_rows_reported,
                      "nodes": len(nodes), "edges": len(edges), "cross_project_exact_titles": cross_edges}, ensure_ascii=False))


if __name__ == "__main__":
    main()
